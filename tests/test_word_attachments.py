from __future__ import annotations

import io
import unittest

from docx import Document

from pdf_attachments import merge_into_prompt
from word_attachments import is_word_bytes, looks_like_word, parse_word


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Колонка A"
    table.rows[0].cells[1].text = "Колонка B"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class LooksLikeWordTests(unittest.TestCase):
    def test_filename_and_mime(self) -> None:
        self.assertTrue(looks_like_word("brief.DOCX", "application/octet-stream"))
        self.assertTrue(
            looks_like_word(
                "x.bin",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        )
        self.assertFalse(looks_like_word("notes.pdf", "application/pdf"))


class ParseWordTests(unittest.TestCase):
    def test_extracts_paragraphs_and_tables(self) -> None:
        data = make_docx_bytes(["Заголовок ТЗ", "Сделай лендинг"])
        self.assertTrue(is_word_bytes(data))
        parsed = parse_word(data, "tz.docx")
        self.assertEqual(parsed.kind, "word")
        self.assertIn("Заголовок ТЗ", parsed.text)
        self.assertIn("Сделай лендинг", parsed.text)
        self.assertIn("Колонка A", parsed.text)
        self.assertIn("Колонка B", parsed.text)

    def test_merges_into_cursor_prompt(self) -> None:
        data = make_docx_bytes(["Пункт первый"])
        parsed = parse_word(data, "spec.docx")
        prompt, images = merge_into_prompt("Внедри это ТЗ", [parsed], [])
        self.assertIn("Внедри это ТЗ", prompt)
        self.assertIn("Word-файл", prompt)
        self.assertIn("Пункт первый", prompt)
        self.assertEqual(images, [])

    def test_rejects_garbage(self) -> None:
        with self.assertRaises(Exception):
            parse_word(b"not a word file", "x.docx")


if __name__ == "__main__":
    unittest.main()

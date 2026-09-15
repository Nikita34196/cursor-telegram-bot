"""Extract text (and a few images) from Word .docx/.doc for Cursor prompts."""

from __future__ import annotations

import io
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document

from pdf_attachments import (
    MAX_IMAGE_BYTES,
    MAX_PDF_BYTES,
    MAX_PDF_TEXT_CHARS,
    MAX_PAGE_IMAGES,
    ParsedPdf,
    PdfAttachmentError,
    _png_payload,
)

WORD_MIMES = frozenset(
    {
        "application/msword",
        "application/vnd.ms-word",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.template",
    }
)
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
DOCX_MAGIC = b"PK\x03\x04"
IMAGE_EXTS = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}


class WordAttachmentError(PdfAttachmentError):
    """User-facing Word processing error."""


def looks_like_word(filename: str | None, mime: str | None) -> bool:
    name = (filename or "").lower()
    mime_type = (mime or "").split(";", 1)[0].strip().lower()
    if mime_type in WORD_MIMES:
        return True
    return name.endswith((".doc", ".docx", ".dotx"))


def is_word_bytes(data: bytes) -> bool:
    if not data:
        return False
    if data.startswith(OLE_MAGIC):
        return True
    if not data.startswith(DOCX_MAGIC):
        return False
    head = data[:32_000]
    return b"wordprocessingml" in head or b"word/document" in head or b"word/" in head


def parse_word(
    data: bytes,
    filename: str,
    *,
    max_images: int = MAX_PAGE_IMAGES,
) -> ParsedPdf:
    if not data:
        raise WordAttachmentError("Word-файл пустой.")
    if len(data) > MAX_PDF_BYTES:
        raise WordAttachmentError(
            f"Word-файл больше {MAX_PDF_BYTES // (1024 * 1024)} МБ — такой файл бот не обрабатывает."
        )

    name = filename or "document.docx"
    if data.startswith(DOCX_MAGIC) or name.lower().endswith(".docx"):
        text, images, warning = _parse_docx(data, max_images)
    elif data.startswith(OLE_MAGIC) or name.lower().endswith(".doc"):
        text, warning = _parse_doc_legacy(data)
        images = []
    else:
        raise WordAttachmentError("Это не Word (.doc/.docx).")

    if not text.strip() and not images:
        raise WordAttachmentError("Не удалось извлечь текст из Word-файла.")

    clipped = text
    if len(clipped) > MAX_PDF_TEXT_CHARS:
        clipped = clipped[:MAX_PDF_TEXT_CHARS] + "\n[текст обрезан]"

    return ParsedPdf(
        filename=name,
        page_count=max(1, clipped.count("\n\n") + 1),
        size_bytes=len(data),
        text=clipped,
        images=images,
        warning=warning,
        kind="word",
    )


def _parse_docx(data: bytes, max_images: int) -> tuple[str, list[dict[str, str]], str | None]:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise WordAttachmentError("Не удалось открыть .docx.") from exc

    parts: list[str] = []
    for paragraph in doc.paragraphs:
        line = (paragraph.text or "").strip()
        if line:
            parts.append(line)
    for table in doc.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [" ".join((cell.text or "").split()) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))

    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                line = (paragraph.text or "").strip()
                if line and line not in parts[:8]:
                    parts.insert(0, line)

    text = "\n".join(parts).strip()
    images = _extract_docx_images(data, max_images)
    warning = None
    if not text:
        warning = "Текста почти нет — если в документе картинки, они переданы как изображения."
    return text, images, warning


def _extract_docx_images(data: bytes, max_images: int) -> list[dict[str, str]]:
    images: list[dict[str, str]] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = [
                n
                for n in archive.namelist()
                if n.startswith("word/media/") and Path(n).suffix.lower() in IMAGE_EXTS
            ]
            for name in names[:max_images]:
                raw = archive.read(name)
                mime = IMAGE_EXTS[Path(name).suffix.lower()]
                if mime == "image/png":
                    payload = _png_payload(raw)
                elif len(raw) <= MAX_IMAGE_BYTES:
                    import base64

                    payload = {
                        "data": base64.b64encode(raw).decode("ascii"),
                        "mimeType": mime,
                    }
                else:
                    payload = None
                if payload:
                    images.append(payload)
    except zipfile.BadZipFile:
        return []
    return images


def _parse_doc_legacy(data: bytes) -> tuple[str, str | None]:
    text = _antiword_text(data)
    if text.strip():
        return text.strip(), None
    text = _utf16_ole_text(data)
    if text.strip():
        return text.strip(), "Старый формат .doc: текст извлечён приблизительно."
    raise WordAttachmentError(
        "Не удалось прочитать .doc. Сохраните файл как .docx и отправьте снова."
    )


def _antiword_text(data: bytes) -> str:
    try:
        with tempfile.NamedTemporaryFile(suffix=".doc") as tmp:
            tmp.write(data)
            tmp.flush()
            result = subprocess.run(
                ["antiword", "-m", "UTF-8.txt", tmp.name],
                capture_output=True,
                timeout=30,
                check=False,
            )
    except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
        return ""
    if result.returncode != 0:
        return ""
    return (result.stdout or b"").decode("utf-8", errors="replace")


def _utf16_ole_text(data: bytes) -> str:
    decoded = data.decode("utf-16le", errors="ignore")
    chunks = re.findall(r"[\wА-Яа-яЁё.,;:!?()\"'«»\-%/ ]{8,}", decoded)
    cleaned: list[str] = []
    for chunk in chunks:
        piece = " ".join(chunk.split())
        if piece and piece not in cleaned:
            cleaned.append(piece)
    return "\n".join(cleaned[:400])
